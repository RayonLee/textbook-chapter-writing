# -*- coding: utf-8 -*-
"""
probe_style.py — 提取样章排版风格（样式名、字号、字体、案例标题格式）
用法: python probe_style.py 样章.docx 起始文字1 起始文字2 ...
  例: python probe_style.py 第一章.docx 第一章 小米智能家居 一、 (一) 1.数据
注意: docx 中可能有修订标记 w:ins，python-docx 段落提取会漏内容，
      提取目录类文档时建议同时直接解析 document.xml 的文本节点。
"""
import re, sys, zipfile
import docx

def main():
    path, anchors = sys.argv[1], sys.argv[2:]
    d = docx.Document(path)
    print('== 样式统计 ==')
    from collections import Counter
    for k, v in Counter(p.style.name for p in d.paragraphs).most_common():
        print(f'  {k}: {v} 段')
    print('== 关键段落格式（样式|字体|字号sz半磅|前40字） ==')
    for p in d.paragraphs:
        t = p.text.strip()
        if not t or not any(t.startswith(a) for a in anchors):
            continue
        xml = re.sub(r'\s+', ' ', re.sub(r' xmlns:[a-z0-9]+="[^"]*"', '', p._element.xml))
        fonts = re.findall(r'w:eastAsia="([^"]+)"', xml)[:2]
        sizes = re.findall(r'<w:sz w:val="(\d+)"/>', xml)[:2]
        print(f'  [{p.style.name}] font={fonts} sz={sizes} | {t[:40]}')

if __name__ == '__main__':
    main()
